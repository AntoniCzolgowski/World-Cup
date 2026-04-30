"""Build a multi-tab XLSX workbook for World Cup outreach mailing lists."""
from __future__ import annotations

import argparse
import csv
import json
import math
import re
import unicodedata
from pathlib import Path
from typing import Any

from enrich_businesses import haversine_m
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
MAILING_DIR = DATA_DIR / "mailing_lists"
CITIES_DIR = DATA_DIR / "cities"
OUTPUT_PATH = MAILING_DIR / "world_cup_outreach.xlsx"
CSV_OUTPUT_PATH = MAILING_DIR / "world_cup_outreach.csv"

DEFAULT_CITIES = ["dallas", "houston", "kansas-city", "monterrey"]
CITY_OVERRIDES: dict[str, dict[str, str]] = {
    "dallas": {"city_label": "Dallas", "tab_name": "Dallas", "unit": "miles"},
    "houston": {"city_label": "Houston", "tab_name": "Houston", "unit": "miles"},
    "kansas-city": {"city_label": "Kansas City", "tab_name": "Kansas City", "unit": "miles"},
    "monterrey": {"city_label": "Monterrey", "tab_name": "Monterrey", "unit": "km"},
}
SIGNATURES = ["Lucas", "Anthony", "Jacob"]

CSV_COLUMNS = [
    "Name",
    "Business_Type",
    "Address",
    "City",
    "Country",
    "Phone",
    "Website",
    "Email",
    "Email_Status",
    "Email_Source",
    "Email_Confidence",
    "Google_Rating",
    "Review_Count",
    "Source",
    "Zone",
]

OUTPUT_COLUMNS = [
    "Email_Sent",
    "Called",
    *CSV_COLUMNS,
    "Distance_To_Stadium",
    "Email_Lucas",
    "Email_Anthony",
    "Email_Jacob",
    "Email_Opened",
]

EMAIL_TEMPLATE = """Hi {business_name} team,

There's a World Cup game {distance_str} from {business_name}, in less than 50 days, and 2\u00d7 {stadium_name} ({venue_2x}) worth of fans coming to {city_label}.

We're CrowdReady - a group of data scientists building actionable reports for {city_label} businesses. We pulled the numbers for {business_name} specifically:

- who's coming (countries, ages, spending profile)
- when they'll be near you, hour by hour, on game day
- a Facebook/Instagram ad you can launch before kickoff\u2014 copy, audience, budget
- a day-of plan: when to staff up, what to stock, what to price up
- the single move with the highest projected revenue for {business_name}

Worth taking a look? I can send over a quick {business_name} snapshot today with the key numbers and the clearest opportunity we see.

If it's useful, 15 minutes on Zoom and I'll show you how to deploy it.

Cheers,
{signature_name}"""


def normalize_key(value: str | None) -> str:
    text = unicodedata.normalize("NFKD", value or "")
    text = text.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def business_match_key(name: str | None, address: str | None) -> str:
    return f"{normalize_key(name)}|{normalize_key(address)}"


def slug_to_label(city: str) -> str:
    return " ".join(part.capitalize() for part in city.split("-"))


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig") as handle:
        return json.load(handle)


def load_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        missing = [column for column in CSV_COLUMNS if column not in (reader.fieldnames or [])]
        if missing:
            raise ValueError(f"{path} is missing required columns: {', '.join(missing)}")
        return [{column: row.get(column, "") for column in CSV_COLUMNS} for row in reader]


def load_stadium_info(city: str) -> dict[str, Any]:
    base_path = CITIES_DIR / city / "base.json"
    base = load_json(base_path)
    stadium = next((node for node in base.get("nodes", []) if node.get("id") == "stadium"), None)
    if not stadium:
        raise ValueError(f"{base_path} does not contain a nodes[].id='stadium' entry")

    city_info = base.get("city", {})
    capacity = int(city_info["venue_capacity"])
    rounded_2x = int(math.ceil((capacity * 2) / 5000) * 5000)
    overrides = CITY_OVERRIDES.get(city, {})
    country = city_info.get("country", "")

    return {
        "stadium_name": city_info.get("venue_name") or stadium.get("label") or slug_to_label(city),
        "stadium_lat": float(stadium["lat"]),
        "stadium_lng": float(stadium["lng"]),
        "venue_2x": f"{rounded_2x // 1000}K",
        "city_label": overrides.get("city_label", slug_to_label(city)),
        "tab_name": overrides.get("tab_name", slug_to_label(city)),
        "unit": overrides.get("unit", "miles" if country == "USA" else "km"),
    }


def load_business_lookup(city: str) -> dict[str, dict[str, Any]]:
    businesses_path = CITIES_DIR / city / "businesses.json"
    payload = load_json(businesses_path)
    lookup: dict[str, dict[str, Any]] = {}
    for business in payload.get("businesses", []):
        key = business_match_key(business.get("name"), business.get("address"))
        if key and key not in lookup:
            lookup[key] = business
    return lookup


def format_distance(distance_km: float, unit: str) -> str:
    if unit == "miles":
        distance = distance_km * 0.621371
        suffix = "miles"
    elif unit == "km":
        distance = distance_km
        suffix = "km"
    else:
        raise ValueError(f"Unsupported distance unit: {unit}")

    if distance < 1:
        value = f"{distance:.1f}"
    elif distance < 10:
        value = f"{distance:.1f}"
    else:
        value = f"{distance:.0f}"
    return f"{value} {suffix}"


def build_email(
    *,
    business_name: str,
    distance_str: str,
    stadium_name: str,
    city_label: str,
    venue_2x: str,
    signature_name: str,
) -> str:
    return EMAIL_TEMPLATE.format(
        business_name=business_name,
        distance_str=distance_str,
        stadium_name=stadium_name,
        venue_2x=venue_2x,
        city_label=city_label,
        signature_name=signature_name,
    )


def available_cities() -> list[str]:
    cities = sorted(path.name.removesuffix("_mailing_list.csv") for path in MAILING_DIR.glob("*_mailing_list.csv"))
    default_first = [city for city in DEFAULT_CITIES if city in cities]
    remaining = [city for city in cities if city not in default_first]
    return [*default_first, *remaining]


def parse_cities(args: argparse.Namespace) -> list[str]:
    if args.all:
        cities = available_cities()
    elif args.cities:
        cities = [city.strip() for city in args.cities.split(",") if city.strip()]
    else:
        cities = DEFAULT_CITIES

    if not cities:
        raise ValueError("No cities selected")
    return cities


def validate_city_inputs(city: str) -> None:
    required_paths = [
        MAILING_DIR / f"{city}_mailing_list.csv",
        CITIES_DIR / city / "base.json",
        CITIES_DIR / city / "businesses.json",
    ]
    missing = [str(path.relative_to(ROOT)) for path in required_paths if not path.exists()]
    if missing:
        raise FileNotFoundError(f"Missing input files for {city}: {', '.join(missing)}")


def append_city_sheet(workbook: Workbook, city: str, is_first_sheet: bool) -> None:
    validate_city_inputs(city)
    rows = load_csv(MAILING_DIR / f"{city}_mailing_list.csv")
    stadium_info = load_stadium_info(city)
    business_lookup = load_business_lookup(city)

    worksheet = workbook.active if is_first_sheet else workbook.create_sheet()
    worksheet.title = stadium_info["tab_name"]
    worksheet.append(OUTPUT_COLUMNS)

    unmatched: list[str] = []
    for csv_row in rows:
        business_key = business_match_key(csv_row["Name"], csv_row["Address"])
        business = business_lookup.get(business_key)
        if not business:
            unmatched.append(csv_row["Name"])
            continue

        distance_m = haversine_m(
            stadium_info["stadium_lat"],
            stadium_info["stadium_lng"],
            float(business["lat"]),
            float(business["lng"]),
        )
        distance_str = format_distance(distance_m / 1000, stadium_info["unit"])
        emails = [
            build_email(
                business_name=csv_row["Name"],
                distance_str=distance_str,
                stadium_name=stadium_info["stadium_name"],
                city_label=stadium_info["city_label"],
                venue_2x=stadium_info["venue_2x"],
                signature_name=signature,
            )
            for signature in SIGNATURES
        ]

        worksheet.append(
            [
                False,
                False,
                *[csv_row[column] for column in CSV_COLUMNS],
                distance_str,
                *emails,
                "",
            ]
        )

    if unmatched:
        examples = ", ".join(unmatched[:5])
        raise ValueError(f"{city}: could not match {len(unmatched)} CSV rows to businesses.json: {examples}")

    format_worksheet(worksheet)


def format_worksheet(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=False, vertical="top")

    email_columns = {"Email_Lucas", "Email_Anthony", "Email_Jacob"}
    validation_columns = {"Email_Sent", "Called", "Email_Opened"}
    header_to_index = {cell.value: cell.column for cell in worksheet[1]}

    validation = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
    worksheet.add_data_validation(validation)
    max_row = max(worksheet.max_row, 2)
    for column_name in validation_columns:
        column_letter = get_column_letter(header_to_index[column_name])
        validation.add(f"{column_letter}2:{column_letter}{max_row}")

    for column_cells in worksheet.columns:
        header = column_cells[0].value
        column_letter = get_column_letter(column_cells[0].column)

        if header in email_columns:
            worksheet.column_dimensions[column_letter].width = 72
            for cell in column_cells[1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            continue

        max_length = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(value), 60))
            if cell.row > 1:
                cell.alignment = Alignment(wrap_text=False, vertical="top")
        worksheet.column_dimensions[column_letter].width = max(12, min(max_length + 2, 62))


def build_workbook(cities: list[str]) -> None:
    workbook = Workbook()
    for index, city in enumerate(cities):
        append_city_sheet(workbook, city, is_first_sheet=index == 0)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_PATH)
    write_combined_csv(workbook)


def csv_cell_value(value: Any) -> str:
    if value is True:
        return "TRUE"
    if value is False:
        return "FALSE"
    if value is None:
        return ""
    return str(value)


def write_combined_csv(workbook: Workbook) -> None:
    with CSV_OUTPUT_PATH.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(OUTPUT_COLUMNS)
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(min_row=2, values_only=True):
                writer.writerow([csv_cell_value(value) for value in row])


def main() -> None:
    parser = argparse.ArgumentParser(description="Build multi-tab XLSX outreach workbook")
    parser.add_argument("--cities", help="Comma-separated city slugs, e.g. dallas,houston,kansas-city")
    parser.add_argument("--all", action="store_true", help="Use all *_mailing_list.csv files found in data/mailing_lists")
    args = parser.parse_args()

    cities = parse_cities(args)
    build_workbook(cities)
    print(f"Wrote {OUTPUT_PATH.relative_to(ROOT)} with {len(cities)} tabs: {', '.join(cities)}")
    print(f"Wrote {CSV_OUTPUT_PATH.relative_to(ROOT)} with combined outreach rows")


if __name__ == "__main__":
    main()
